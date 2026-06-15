import os
from django.core.management.base import BaseCommand
from django.core.mail import EmailMultiAlternatives
from django.conf import settings
import openpyxl
from docx import Document
from docx2pdf import convert
import tempfile


class Command(BaseCommand):
    help = 'Genera y envía cartas de invitación personalizadas desde el Excel'

    def add_arguments(self, parser):
        parser.add_argument('--file', type=str, required=True)
        parser.add_argument('--plantillas', type=str, default='docs/correos/plantillas')
        parser.add_argument('--dry-run', action='store_true')

    def handle(self, *args, **options):
        excel_path = options['file']
        plantillas_dir = options['plantillas']
        dry_run = options['dry_run']

        config = {
            'colegio': {
                'plantilla': 'invitacion_colegios.docx',
                'marcador': '{{ NOMBRE DE COLEGIO }}',
            },
            'universidad': {
                'plantilla': 'invitacion_universidades.docx',
                'marcador': '{{NOMBRE UNIVERSIDAD}}',
            },
            'patrocinador': {
                'plantilla': 'invitacion_patrocinadores.docx',
                'marcador': '{{NOMBRE}}',
            },
        }

        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        headers = [cell.value for cell in ws[1]]
        col = {h.strip(): i for i, h in enumerate(headers) if h}

        if 'ESTADO' not in col:
            self.stdout.write(self.style.ERROR('❌ No se encontró la columna ESTADO en el Excel'))
            return

        enviados = 0
        omitidos = 0

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            nombre = row[col['NOMBRE']]
            correo = row[col['CORREO']]
            tipo   = str(row[col['TIPO']]).lower().strip() if row[col['TIPO']] else ''
            estado = str(row[col['ESTADO']]).lower().strip() if row[col['ESTADO']] else ''

            if not nombre or not correo or tipo not in config:
                self.stdout.write(f'  ⚠ Fila omitida: {nombre} / {tipo}')
                continue

            if estado == 'enviado':
                self.stdout.write(f'  ⏭ Ya enviado: {nombre}')
                omitidos += 1
                continue

            cfg = config[tipo]
            plantillas_dir_abs = os.path.normpath(os.path.join(settings.BASE_DIR, plantillas_dir))
            plantilla_path = os.path.normpath(os.path.join(plantillas_dir_abs, cfg['plantilla']))

            doc = Document(plantilla_path)
            self._reemplazar(doc, cfg['marcador'], nombre)

            with tempfile.TemporaryDirectory() as tmpdir:
                docx_out = os.path.join(tmpdir, f'carta_{nombre}.docx')
                pdf_out  = os.path.join(tmpdir, f'carta_{nombre}.pdf')
                doc.save(docx_out)
                convert(docx_out, pdf_out)

                if dry_run:
                    self.stdout.write(f'  [DRY-RUN] {nombre} → {correo} ({tipo})')
                else:
                    self._enviar(correo, nombre, tipo, pdf_out)
                    ws.cell(row=row_idx, column=col['ESTADO'] + 1).value = 'enviado'
                    self.stdout.write(self.style.SUCCESS(f'  ✓ Enviado: {nombre} → {correo}'))
                    enviados += 1

        if not dry_run:
            wb.save(excel_path)
            self.stdout.write(self.style.SUCCESS('✅ Excel actualizado'))

        self.stdout.write(f'\nTotal enviados: {enviados} | Ya enviados antes: {omitidos}')

    def _reemplazar(self, doc, marcador, nombre):
        for para in doc.paragraphs:
            if marcador in para.text:
                texto_completo = para.text.replace(marcador, nombre)
                for i, run in enumerate(para.runs):
                    run.text = texto_completo if i == 0 else ''
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if marcador in para.text:
                            texto_completo = para.text.replace(marcador, nombre)
                            for i, run in enumerate(para.runs):
                                run.text = texto_completo if i == 0 else ''

    def _enviar(self, correo, nombre, tipo, pdf_path):
        asuntos = {
            'colegio':      'Invitación - Semana de la Electrónica UTP 2026',
            'universidad':  'Invitación - Semana de la Electrónica UTP 2026',
            'patrocinador': 'Invitación a patrocinar - Semana de la Electrónica UTP 2026',
        }
        cuerpo_html = f"""
        <p>Estimados señores de <strong>{nombre}</strong>,</p>
        <p>Adjunto encontrarán la carta de invitación.</p>
        <br>
        <p>Atentamente,<br>Programa de Ingeniería Electrónica - UTP</p>
        """
        email = EmailMultiAlternatives(
            subject=asuntos[tipo],
            body=f'Estimados señores de {nombre}, adjunto encontrarán la carta de invitación.\n\nAtentamente,\nPrograma de Ingeniería Electrónica - UTP',
            from_email=settings.DEFAULT_FROM_EMAIL,
            to=[correo],
        )
        email.attach_alternative(cuerpo_html, "text/html")
        with open(pdf_path, 'rb') as f:
            email.attach(f'Invitacion_{nombre}.pdf', f.read(), 'application/pdf')
        email.send()