import os
import re
import tempfile
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.mail import EmailMultiAlternatives
from django.core.validators import validate_email
from django.utils import timezone
from docx import Document
from openpyxl import load_workbook

from .models import (
    CommunicationBatch,
    CommunicationLog,
    CommunicationLogStatus,
    CommunicationSendMode,
    CommunicationStatus,
    CommunicationTemplate,
)


MARKER_RE = re.compile(r"\{\{\s*[^{}]+\s*\}\}")
SMTP_BACKEND = "django.core.mail.backends.smtp.EmailBackend"
SMTP_BLOCKED_MESSAGE = "No hay configuracion SMTP completa. El envio real esta bloqueado por seguridad."
DEFAULT_EVENT_NAME = "Torneo Robot Explota Globos UTP"


@dataclass(frozen=True)
class EmailConfigurationStatus:
    can_send_real: bool
    backend: str
    message: str


def marker_name(marker: str) -> str:
    return marker.strip()[2:-2].strip()


def normalize_marker(marker: str) -> str:
    return "{{" + marker_name(marker).upper() + "}}"


def parse_required_markers(value):
    if isinstance(value, list):
        raw_markers = value
    else:
        raw_markers = re.split(r"[\n,;]+", value or "")
    markers = []
    for item in raw_markers:
        marker = str(item).strip()
        if not marker:
            continue
        if not marker.startswith("{{"):
            marker = "{{" + marker
        if not marker.endswith("}}"):
            marker = marker + "}}"
        normalized = normalize_marker(marker)
        if normalized not in markers:
            markers.append(normalized)
    return markers


def _file_path(file_or_path):
    if hasattr(file_or_path, "path"):
        return file_or_path.path
    return str(file_or_path)


def _open_document(file_or_path):
    if hasattr(file_or_path, "path"):
        return Document(file_or_path.path)
    if hasattr(file_or_path, "seek"):
        file_or_path.seek(0)
        return Document(file_or_path)
    return Document(str(file_or_path))


def _docx_text_parts(doc):
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return parts


def extract_docx_text_markers(file_or_path):
    doc = _open_document(file_or_path)
    text = "\n".join(_docx_text_parts(doc))
    return sorted({normalize_marker(match.group(0)) for match in MARKER_RE.finditer(text)})


def validate_template_markers(template_or_file, required_markers):
    required = parse_required_markers(required_markers)
    found = set(extract_docx_text_markers(template_or_file.file if hasattr(template_or_file, "file") else template_or_file))
    missing = [marker for marker in required if marker not in found]
    return missing


def sanitize_filename(value):
    text = unicodedata.normalize("NFKD", str(value or "archivo")).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("._")
    return text[:80] or "archivo"


def validate_email_address(email):
    try:
        validate_email(email)
    except ValidationError:
        return False
    return True


def _context_lookup(context, marker):
    name = marker_name(marker)
    candidates = [
        marker,
        normalize_marker(marker),
        name,
        name.upper(),
        name.lower(),
    ]
    for key in candidates:
        if key in context and context[key] is not None:
            return str(context[key])
    return ""


def _replace_in_paragraph(paragraph, context):
    original = paragraph.text
    if not MARKER_RE.search(original):
        return
    rendered = original
    for match in MARKER_RE.finditer(original):
        rendered = rendered.replace(match.group(0), _context_lookup(context, match.group(0)))
    for index, run in enumerate(paragraph.runs):
        run.text = rendered if index == 0 else ""


def render_docx_template(template_file, context, output_path=None):
    doc = _open_document(template_file)
    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, context)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, context)

    if output_path is None:
        tmpdir = tempfile.mkdtemp(prefix="communication_")
        output_path = os.path.join(tmpdir, "rendered.docx")
    doc.save(output_path)
    return output_path


def build_invitation_context(recipient, extra_data=None):
    extra_data = extra_data or {}
    if isinstance(recipient, dict):
        name = recipient.get("name") or recipient.get("NOMBRE") or ""
        institution = recipient.get("institution_name") or recipient.get("INSTITUCION") or name
        email = recipient.get("email") or recipient.get("CORREO") or ""
    else:
        name = getattr(recipient, "name", "")
        institution = getattr(recipient, "institution_name", "") or name
        email = getattr(recipient, "email", "")

    context = {
        "NOMBRE": name,
        "INSTITUCION": institution,
        "CORREO": email,
        "EVENTO": extra_data.get("EVENTO", DEFAULT_EVENT_NAME),
        "FECHA": extra_data.get("FECHA", timezone.localdate().isoformat()),
        "ENLACE_CONFIRMACION": extra_data.get("ENLACE_CONFIRMACION", ""),
    }
    for key, value in extra_data.items():
        context[str(key).upper()] = value
    return context


def is_real_email_configured():
    backend = getattr(settings, "EMAIL_BACKEND", "")
    if backend != SMTP_BACKEND:
        return EmailConfigurationStatus(
            can_send_real=False,
            backend=backend,
            message="Backend de correo en modo consola/desarrollo. El envio real esta bloqueado.",
        )

    required_values = {
        "EMAIL_HOST": getattr(settings, "EMAIL_HOST", ""),
        "EMAIL_PORT": getattr(settings, "EMAIL_PORT", ""),
        "EMAIL_HOST_USER": getattr(settings, "EMAIL_HOST_USER", ""),
        "EMAIL_HOST_PASSWORD": getattr(settings, "EMAIL_HOST_PASSWORD", ""),
        "DEFAULT_FROM_EMAIL": getattr(settings, "DEFAULT_FROM_EMAIL", ""),
    }
    missing = [name for name, value in required_values.items() if not str(value or "").strip()]
    if missing:
        return EmailConfigurationStatus(
            can_send_real=False,
            backend=backend,
            message=f"{SMTP_BLOCKED_MESSAGE} Faltan: {', '.join(missing)}.",
        )
    return EmailConfigurationStatus(
        can_send_real=True,
        backend=backend,
        message="SMTP configurado. El envio real requiere confirmacion explicita.",
    )


def create_communication_batch(
    name,
    communication_type,
    template=None,
    dry_run=True,
    created_by=None,
    send_mode=None,
    test_recipient="",
):
    send_mode = send_mode or (CommunicationSendMode.DRY_RUN if dry_run else CommunicationSendMode.OFFICIAL)
    return CommunicationBatch.objects.create(
        name=name,
        communication_type=communication_type,
        template=template,
        dry_run=dry_run,
        send_mode=send_mode,
        test_recipient=test_recipient or "",
        created_by=created_by if getattr(created_by, "is_authenticated", False) else None,
        status=CommunicationStatus.PREVIEWED if dry_run else CommunicationStatus.SENDING,
    )


def _create_or_update_log(
    *,
    log=None,
    batch=None,
    recipient_name,
    recipient_email,
    original_recipient_email="",
    physical_recipient_email="",
    communication_type,
    subject,
    status,
    error_message="",
    rendered_preview="",
    sent_at=None,
):
    if log is None:
        return CommunicationLog.objects.create(
            batch=batch,
            recipient_name=recipient_name,
            recipient_email=recipient_email or "",
            original_recipient_email=original_recipient_email or recipient_email or "",
            physical_recipient_email=physical_recipient_email or recipient_email or "",
            communication_type=communication_type,
            subject=subject,
            status=status,
            error_message=error_message,
            rendered_preview=rendered_preview[:4000],
            sent_at=sent_at,
        )
    log.status = status
    log.error_message = error_message
    log.rendered_preview = rendered_preview[:4000]
    log.sent_at = sent_at
    if original_recipient_email:
        log.original_recipient_email = original_recipient_email
    if physical_recipient_email:
        log.physical_recipient_email = physical_recipient_email
    log.save(
        update_fields=[
            "status",
            "error_message",
            "rendered_preview",
            "sent_at",
            "original_recipient_email",
            "physical_recipient_email",
            "updated_at",
        ]
    )
    return log


def send_communication_email(
    *,
    batch=None,
    recipient_name,
    recipient_email,
    communication_type,
    subject,
    body,
    html_body="",
    attachments=None,
    dry_run=True,
    test_recipient="",
    rendered_preview="",
):
    attachments = attachments or []
    if not validate_email_address(recipient_email):
        return _create_or_update_log(
            batch=batch,
            recipient_name=recipient_name,
            recipient_email=recipient_email,
            original_recipient_email=recipient_email,
            communication_type=communication_type,
            subject=subject,
            status=CommunicationLogStatus.SKIPPED,
            error_message="Correo invalido.",
            rendered_preview=rendered_preview,
        )

    if dry_run:
        return _create_or_update_log(
            batch=batch,
            recipient_name=recipient_name,
            recipient_email=recipient_email,
            original_recipient_email=recipient_email,
            communication_type=communication_type,
            subject=subject,
            status=CommunicationLogStatus.DRY_RUN,
            rendered_preview=rendered_preview,
        )

    controlled_test = bool(test_recipient)
    physical_recipient = test_recipient if controlled_test else recipient_email
    if controlled_test and not validate_email_address(test_recipient):
        return _create_or_update_log(
            batch=batch,
            recipient_name=recipient_name,
            recipient_email=recipient_email,
            original_recipient_email=recipient_email,
            physical_recipient_email=test_recipient,
            communication_type=communication_type,
            subject=subject,
            status=CommunicationLogStatus.FAILED,
            error_message="Correo de prueba invalido.",
            rendered_preview=rendered_preview,
        )

    email_status = is_real_email_configured()
    if not email_status.can_send_real:
        return _create_or_update_log(
            batch=batch,
            recipient_name=recipient_name,
            recipient_email=recipient_email,
            original_recipient_email=recipient_email,
            physical_recipient_email=physical_recipient,
            communication_type=communication_type,
            subject=subject,
            status=CommunicationLogStatus.FAILED,
            error_message=email_status.message,
            rendered_preview=rendered_preview,
        )

    try:
        final_subject = f"[PRUEBA CONTROLADA] {subject}" if controlled_test else subject
        final_body = (
            "PRUEBA CONTROLADA: este correo se envio al destinatario de prueba, no al destinatario original.\n"
            f"Destinatario original: {recipient_name} <{recipient_email}>\n\n"
            f"{body}"
            if controlled_test
            else body
        )
        final_html = (
            f"<p><strong>PRUEBA CONTROLADA:</strong> este correo se envio al destinatario de prueba, "
            f"no al destinatario original.</p><p>Destinatario original: {recipient_name} &lt;{recipient_email}&gt;</p>"
            f"{html_body}"
            if controlled_test and html_body
            else html_body
        )
        email = EmailMultiAlternatives(
            subject=final_subject,
            body=final_body,
            from_email=settings.DEFAULT_FROM_EMAIL,
            to=[physical_recipient],
        )
        if final_html:
            email.attach_alternative(final_html, "text/html")
        for attachment in attachments:
            email.attach(*attachment)
        email.send(fail_silently=False)
    except Exception as error:
        return _create_or_update_log(
            batch=batch,
            recipient_name=recipient_name,
            recipient_email=recipient_email,
            original_recipient_email=recipient_email,
            physical_recipient_email=physical_recipient,
            communication_type=communication_type,
            subject=subject,
            status=CommunicationLogStatus.FAILED,
            error_message=str(error),
            rendered_preview=rendered_preview,
        )

    return _create_or_update_log(
        batch=batch,
        recipient_name=recipient_name,
        recipient_email=recipient_email,
        original_recipient_email=recipient_email,
        physical_recipient_email=physical_recipient,
        communication_type=communication_type,
        subject=subject,
        status=CommunicationLogStatus.TEST_SENT if controlled_test else CommunicationLogStatus.SENT,
        rendered_preview=rendered_preview,
        sent_at=timezone.now(),
    )


def mark_batch_finished(batch):
    statuses = list(batch.logs.values_list("status", flat=True))
    if not statuses:
        batch.status = CommunicationStatus.CANCELLED
    elif any(status == CommunicationLogStatus.FAILED for status in statuses):
        batch.status = CommunicationStatus.FAILED
    elif all(status == CommunicationLogStatus.DRY_RUN for status in statuses):
        batch.status = CommunicationStatus.PREVIEWED
    elif any(status == CommunicationLogStatus.SENT for status in statuses):
        batch.status = CommunicationStatus.SENT
        batch.sent_at = timezone.now()
    elif any(status == CommunicationLogStatus.TEST_SENT for status in statuses):
        batch.status = CommunicationStatus.SENT
        batch.sent_at = timezone.now()
    else:
        batch.status = CommunicationStatus.CANCELLED
    batch.save(update_fields=["status", "sent_at", "updated_at"])
    return batch


def active_template_for(template_type):
    return CommunicationTemplate.objects.filter(template_type=template_type, is_active=True).order_by("-updated_at").first()


def load_recipients_from_xlsx(file_or_path):
    wb = load_workbook(file_or_path, data_only=True, read_only=True)
    ws = wb.active
    headers = [str(cell.value or "").strip().upper() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    column = {header: index for index, header in enumerate(headers) if header}
    required = {"NOMBRE", "CORREO", "TIPO"}
    missing = sorted(required - set(column))
    if missing:
        raise ValueError("Faltan columnas requeridas en el Excel: " + ", ".join(missing))

    recipients = []
    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None for value in row):
            continue
        recipient = {
            "row": row_index,
            "name": row[column["NOMBRE"]] or "",
            "email": row[column["CORREO"]] or "",
            "recipient_type": str(row[column["TIPO"]] or "").strip().lower(),
            "institution_name": row[column["INSTITUCION"]] if "INSTITUCION" in column else (row[column["NOMBRE"]] or ""),
            "extra_data": {},
        }
        for header, index in column.items():
            if header not in {"NOMBRE", "CORREO", "TIPO", "INSTITUCION"} and index < len(row):
                recipient["extra_data"][header] = row[index]
        recipients.append(recipient)
    return recipients


def read_attachment_bytes(path):
    path = Path(path)
    return (path.name, path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def invitation_subject(communication_type):
    subjects = {
        "colegio": "Invitacion - Semana de la Electronica UTP",
        "universidad": "Invitacion - Semana de la Electronica UTP",
        "patrocinador": "Invitacion a patrocinar - Semana de la Electronica UTP",
        "general": "Comunicacion - Robot Explota Globos UTP",
    }
    return subjects.get(communication_type, subjects["general"])


def send_rendered_invitation(
    *, batch, template, recipient, communication_type, dry_run=True, extra_data=None, test_recipient=""
):
    context = build_invitation_context(recipient, extra_data=extra_data)
    recipient_name = context["NOMBRE"]
    recipient_email = context["CORREO"]
    subject = invitation_subject(communication_type)
    body = (
        f"Hola,\n\n"
        f"Adjuntamos la invitacion personalizada para {recipient_name}.\n\n"
        f"{DEFAULT_EVENT_NAME}\n"
        f"Universidad Tecnologica de Pereira"
    )
    html_body = (
        f"<p>Hola,</p><p>Adjuntamos la invitacion personalizada para "
        f"<strong>{recipient_name}</strong>.</p><p>{DEFAULT_EVENT_NAME}<br>Universidad Tecnologica de Pereira</p>"
    )
    rendered_preview = "\n".join(f"{key}: {value}" for key, value in context.items() if value)[:1200]

    attachments = []
    try:
        with tempfile.TemporaryDirectory(prefix="communication_invitation_") as tmpdir:
            filename = f"Invitacion_{sanitize_filename(recipient_name)}.docx"
            output_path = os.path.join(tmpdir, filename)
            render_docx_template(template.file, context, output_path=output_path)
            attachments.append(read_attachment_bytes(output_path))
            return send_communication_email(
                batch=batch,
                recipient_name=recipient_name,
                recipient_email=recipient_email,
                communication_type=communication_type,
                subject=subject,
                body=body,
                html_body=html_body,
                attachments=attachments,
                dry_run=dry_run,
                test_recipient=test_recipient,
                rendered_preview=rendered_preview,
            )
    except Exception as error:
        return _create_or_update_log(
            batch=batch,
            recipient_name=recipient_name,
            recipient_email=recipient_email,
            communication_type=communication_type,
            subject=subject,
            status=CommunicationLogStatus.FAILED,
            error_message=str(error),
            rendered_preview=rendered_preview,
        )


def send_attendance_request_for_registration(
    *, batch, registration, registration_type, confirmation_url, dry_run=True, test_recipient=""
):
    subject = "Confirma tu asistencia - Torneo Robot Explota Globos"
    body = (
        f"Hola,\n\n"
        f"Ya estamos en la etapa de confirmacion de asistencia para el torneo.\n\n"
        f"Robot: {registration.robot_name}\n"
        f"Categoria: {registration_type}\n"
        f"Institucion: {registration.institution_name}\n\n"
        f"Para confirmar que asistiras al evento, entra al siguiente enlace:\n"
        f"{confirmation_url}\n\n"
        f"Solo los robots que confirmen asistencia seran tenidos en cuenta para generar el cuadro oficial.\n\n"
        f"Torneo Robot Explota Globos\n"
        f"Universidad Tecnologica de Pereira"
    )
    html_body = (
        f"<p>Hola,</p><p>Ya estamos en la etapa de confirmacion de asistencia para el torneo.</p>"
        f"<p><strong>Robot:</strong> {registration.robot_name}<br>"
        f"<strong>Categoria:</strong> {registration_type}<br>"
        f"<strong>Institucion:</strong> {registration.institution_name}</p>"
        f"<p><a href=\"{confirmation_url}\">Confirmar asistencia</a></p>"
    )
    log = send_communication_email(
        batch=batch,
        recipient_name=f"{registration.institution_name} - {registration.robot_name}",
        recipient_email=registration.contact_email,
        communication_type="asistencia",
        subject=subject,
        body=body,
        html_body=html_body,
        dry_run=dry_run,
        test_recipient=test_recipient,
        rendered_preview=confirmation_url,
    )
    if log.status == CommunicationLogStatus.SENT:
        registration.attendance_request_sent_at = timezone.now()
        registration.save(update_fields=["attendance_request_sent_at", "updated_at"])
    return log
