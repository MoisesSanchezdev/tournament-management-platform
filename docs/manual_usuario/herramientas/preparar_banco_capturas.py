import json
import os
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "docs.manual_usuario.herramientas.demo_settings")

import django  # noqa: E402
from django.conf import settings  # noqa: E402
from django.contrib.auth import get_user_model  # noqa: E402
from django.db import connection, transaction  # noqa: E402
from django.utils import timezone  # noqa: E402
from docx import Document  # noqa: E402


def assert_demo_database() -> None:
    db_path = Path(settings.DATABASES["default"]["NAME"]).resolve()
    demo_root = settings.DEMO_ROOT.resolve()
    official_db = (settings.BASE_DIR / "db.sqlite3").resolve()
    if demo_root not in db_path.parents:
        raise SystemExit(f"La base no esta dentro del entorno demo: {db_path}")
    if db_path == official_db:
        raise SystemExit("La base demo apunta a la base oficial. Abortado.")


def create_docx_template(relative_name: str) -> str:
    target_dir = Path(settings.MEDIA_ROOT) / "communication_templates"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / relative_name
    doc = Document()
    doc.add_heading("Invitacion demo Robot Explota Globos", level=1)
    doc.add_paragraph("Hola {{NOMBRE}}, invitamos a {{INSTITUCION}} al {{EVENTO}}.")
    doc.add_paragraph("Correo ficticio: {{CORREO}}")
    doc.save(target)
    return f"communication_templates/{relative_name}"


def safe_clean():
    from apps.invitations.models import CommunicationBatch, CommunicationLog, CommunicationRecipient, CommunicationTemplate
    from apps.participants.models import Institution, SchoolParticipant, SchoolRegistration, Team, TeamMember, UniversityParticipant, UniversityRegistration
    from apps.tournament.models import (
        CompetitionBattle,
        CompetitionBattleEntry,
        CompetitionHistoryEntry,
        DivisionCompetition,
        DivisionGroup,
        DivisionGroupEntry,
        Match,
        RuleSection,
        TeamCompetitionState,
        TournamentEdition,
        TournamentPhase,
    )

    User = get_user_model()
    CommunicationLog.objects.all().delete()
    CommunicationBatch.objects.all().delete()
    CommunicationRecipient.objects.all().delete()
    CommunicationTemplate.objects.all().delete()
    CompetitionHistoryEntry.objects.all().delete()
    TeamCompetitionState.objects.all().delete()
    CompetitionBattleEntry.objects.all().delete()
    CompetitionBattle.objects.all().delete()
    DivisionGroupEntry.objects.all().delete()
    DivisionGroup.objects.all().delete()
    DivisionCompetition.objects.all().delete()
    Match.objects.all().delete()
    TournamentPhase.objects.all().delete()
    TeamMember.objects.all().delete()
    Team.objects.all().delete()
    Institution.objects.all().delete()
    SchoolParticipant.objects.all().delete()
    UniversityParticipant.objects.all().delete()
    SchoolRegistration.objects.all().delete()
    UniversityRegistration.objects.all().delete()
    RuleSection.objects.all().delete()
    TournamentEdition.objects.all().delete()
    User.objects.filter(username="operador_documentacion").delete()


def main() -> None:
    django.setup()
    assert_demo_database()

    from apps.invitations.models import (
        CommunicationBatch,
        CommunicationLog,
        CommunicationLogStatus,
        CommunicationRecipient,
        CommunicationSendMode,
        CommunicationStatus,
        CommunicationTemplate,
    )
    from apps.participants.models import (
        Institution,
        InstitutionType,
        RegistrationStatus,
        SchoolParticipant,
        SchoolRegistration,
        Team,
        TeamMember,
        UniversityParticipant,
        UniversityRegistration,
    )
    from apps.tournament.models import (
        BattleFormat,
        CompetitionBattle,
        CompetitionBattleEntry,
        CompetitionHistoryEntry,
        CompetitionStage,
        DivisionCompetition,
        MatchStatus,
        ParticipantStatus,
        RuleSection,
        TeamCompetitionState,
        TournamentEdition,
    )
    from apps.tournament.services import (
        initialize_competition,
        materialize_recommended_duel_stage,
        materialize_repechage_stage,
        record_history,
        save_final_podium,
        set_battle_qualifiers,
        set_battle_winner,
        set_group_qualifiers,
    )

    User = get_user_model()
    with transaction.atomic():
        safe_clean()
        user = User.objects.create_user(
            username="operador_documentacion",
            email="operador@example.com",
            password="DemoManualUsuario2026!",
            first_name="Operador",
            last_name="Documentacion",
            is_staff=True,
            is_superuser=True,
        )

        edition = TournamentEdition.objects.create(
            name="Edicion Demo Manual de Usuario",
            description="Edicion ficticia y aislada para banco definitivo de capturas base.",
            location="Laboratorio Demo UTP",
            is_active=True,
        )
        rules = [
            ("Registro demo", "Cada robot demo conserva datos ficticios y correos example.com."),
            ("Formato competitivo", "La competencia demo muestra grupos, batallas, repechaje y podio."),
            ("Comunicaciones seguras", "El entorno demo usa backend de correo local y no envia SMTP."),
            ("Privacidad", "Las capturas base no deben exponer tokens, documentos ni credenciales."),
        ]
        for order, (title, summary) in enumerate(rules, start=1):
            RuleSection.objects.create(edition=edition, title=title, summary=summary, body=summary, order=order, is_published=True)

        school_institutions = [
            Institution.objects.create(name="Institucion Demo Norte", institution_type=InstitutionType.SCHOOL, city="Ciudad Demo"),
            Institution.objects.create(name="Colegio Demo Sur", institution_type=InstitutionType.SCHOOL, city="Ciudad Demo"),
            Institution.objects.create(name="Instituto Demo Oriente", institution_type=InstitutionType.SCHOOL, city="Ciudad Demo"),
            Institution.objects.create(name="Liceo Demo Occidente", institution_type=InstitutionType.SCHOOL, city="Ciudad Demo"),
        ]
        university = Institution.objects.create(name="Universidad de Ejemplo", institution_type=InstitutionType.UNIVERSITY, city="Ciudad Demo")

        school_names = [
            "Equipo Andromeda", "Equipo Vector", "Equipo Kilobyte", "Equipo Prisma",
            "Equipo Nebula", "Equipo Pixel", "Equipo Orion", "Equipo Lambda",
            "Equipo Sigma", "Equipo Delta", "Equipo Quasar", "Equipo Atlas",
            "Equipo Nova", "Equipo Zenith", "Equipo Cobalto", "Equipo Aurora",
        ]
        university_names = [
            "Equipo Byte", "Equipo Kernel", "Equipo Quantum", "Equipo Circuito",
            "Equipo Pulso", "Equipo Magneto", "Equipo Cosmos", "Equipo Servo",
        ]

        created_teams = []
        for index, name in enumerate(school_names, start=1):
            institution = school_institutions[(index - 1) % len(school_institutions)]
            team = Team.objects.create(
                edition=edition,
                name=name,
                institution=institution,
                category_label="Colegios",
                robot_name=name.replace("Equipo ", "Robot "),
                coach_name=f"Responsable Demo {index:02d}",
                coach_email=f"responsable{index:02d}@example.com",
                coach_phone=f"3000000{index:03d}",
                status=RegistrationStatus.APPROVED,
                notes="Equipo ficticio para capturas del Manual de Usuario.",
            )
            TeamMember.objects.create(team=team, full_name=f"Participante Demo {index:02d}", document_number=f"900000{index:04d}", email=f"participante{index:02d}@example.com", age=15, role="lider", is_team_lead=True)
            created_teams.append(team)

        for offset, name in enumerate(university_names, start=1):
            index = offset + len(created_teams)
            team = Team.objects.create(
                edition=edition,
                name=name,
                institution=university,
                category_label="Universidades",
                robot_name=name.replace("Equipo ", "Robot "),
                coach_name=f"Responsable Universidad Demo {offset:02d}",
                coach_email=f"universidad{offset:02d}@example.com",
                coach_phone=f"3010000{offset:03d}",
                status=RegistrationStatus.APPROVED,
                notes="Equipo universitario ficticio para capturas.",
            )
            TeamMember.objects.create(team=team, full_name=f"Participante Universidad Demo {offset:02d}", document_number=f"920000{offset:04d}", email=f"universitario{offset:02d}@example.com", age=19, role="lider", is_team_lead=True)

        # Public registration examples.
        pending_school = SchoolRegistration.objects.create(
            edition=edition,
            institution_name="Institucion Registro Demo",
            responsible_name="Responsable Registro Demo",
            robot_name="Robot Registro Demo",
            contact_phone="3110000001",
            contact_email="registro.demo@example.com",
            status=RegistrationStatus.SUBMITTED,
        )
        SchoolParticipant.objects.create(registration=pending_school, full_name="Participante Registro Demo", document_number="9100000001", role="lider", is_team_lead=True)
        confirmed_school = SchoolRegistration.objects.create(
            edition=edition,
            institution_name="Institucion Confirmada Demo",
            responsible_name="Responsable Confirmado Demo",
            robot_name="Robot Confirmado Demo",
            contact_phone="3110000002",
            contact_email="confirmado.demo@example.com",
            status=RegistrationStatus.APPROVED,
            attendance_confirmed_at=timezone.now(),
        )
        SchoolParticipant.objects.create(registration=confirmed_school, full_name="Participante Confirmado Demo", document_number="9100000002", role="lider", is_team_lead=True)
        university_reg = UniversityRegistration.objects.create(
            edition=edition,
            institution_name="Universidad Registro Demo",
            responsible_name="Responsable Universidad Demo",
            robot_name="Robot Universidad Demo",
            semester=3,
            contact_phone="3120000001",
            contact_email="universidad.demo@example.com",
            status=RegistrationStatus.SUBMITTED,
        )
        UniversityParticipant.objects.create(registration=university_reg, full_name="Participante Universidad Demo", document_number="9300000001", role="lider", is_team_lead=True)

        school_competition = initialize_competition(edition, "school")
        university_competition = initialize_competition(edition, "university")

        # Close all school groups to make recommended phase available.
        for group in school_competition.groups.prefetch_related("entries").order_by("order"):
            selected = [str(entry.id) for entry in group.entries.order_by("slot_order")[:2]]
            set_group_qualifiers(group, selected, manual_override=False)

        # Create deterministic semifinal and final fixtures for documentation.
        qualified_teams = [
            entry.team
            for group in school_competition.groups.prefetch_related("entries__team").order_by("order")
            for entry in group.entries.order_by("slot_order")
            if entry.qualified_from_group
        ]
        semifinal_battles = []
        for battle_order in range(2):
            battle = CompetitionBattle.objects.create(
                competition=school_competition,
                stage=CompetitionStage.SEMIFINAL,
                format_type=BattleFormat.DUEL,
                order=battle_order + 1,
                name=f"Semifinal demo {battle_order + 1}",
                status=MatchStatus.PENDING,
            )
            semifinal_battles.append(battle)
            for slot, team in enumerate(qualified_teams[battle_order * 2 : battle_order * 2 + 2], start=1):
                CompetitionBattleEntry.objects.create(battle=battle, team=team, slot_order=slot, origin_label="Clasificado demo")
                state = TeamCompetitionState.objects.get(competition=school_competition, team=team)
                state.current_stage = CompetitionStage.SEMIFINAL
                state.current_battle = battle
                state.current_status = ParticipantStatus.ACTIVE
                state.save(update_fields=["current_stage", "current_battle", "current_status", "updated_at"])
        semifinal_winners = []
        for battle in semifinal_battles:
            first_entry = battle.entries.order_by("slot_order").first()
            if first_entry:
                set_battle_winner(battle, first_entry.team_id, manual_override=False)
                semifinal_winners.append(first_entry.team)

        final_battle = CompetitionBattle.objects.create(
            competition=school_competition,
            stage=CompetitionStage.FINAL,
            format_type=BattleFormat.DUEL,
            order=1,
            name="Gran final demo",
            status=MatchStatus.PENDING,
        )
        for slot, team in enumerate(semifinal_winners[:2], start=1):
            CompetitionBattleEntry.objects.create(battle=final_battle, team=team, slot_order=slot, origin_label="Ganador semifinal demo")
            state = TeamCompetitionState.objects.get(competition=school_competition, team=team)
            state.current_stage = CompetitionStage.FINAL
            state.current_battle = final_battle
            state.current_status = ParticipantStatus.ACTIVE
            state.save(update_fields=["current_stage", "current_battle", "current_status", "updated_at"])
        final_entries = list(final_battle.entries.order_by("slot_order"))
        if final_entries:
            set_battle_winner(final_battle, final_entries[0].team_id, manual_override=False)
        if len(final_entries) >= 2:
            save_final_podium(
                school_competition,
                champion_team_id=final_entries[0].team_id,
                second_team_id=final_entries[1].team_id,
                third_team_id=None,
            )
        # Add repechage from eliminated candidates.
        try:
            materialize_repechage_stage(school_competition)
        except Exception:
            pass

        # Add a custom multi-qualifier battle on quarterfinal for visual coverage.
        configuration = dict(school_competition.configuration or {})
        stages = list(configuration.get("stages") or [])
        multi_stage = {
            "stage": CompetitionStage.QUARTERFINAL,
            "count": 1,
            "title": "Batalla de clasificados demo",
            "multi_qualifier_enabled": True,
            "qualifier_targets": [2],
            "qualifiers_per_battle": 2,
            "total_qualifiers": 2,
            "kind": "manual_demo",
        }
        for item in stages:
            if item.get("stage") == CompetitionStage.QUARTERFINAL:
                item.update(multi_stage)
                break
        else:
            stages.append(multi_stage)
        configuration["stages"] = stages
        school_competition.configuration = configuration
        school_competition.save(update_fields=["configuration", "updated_at"])
        multi_battle = CompetitionBattle.objects.create(
            competition=school_competition,
            stage=CompetitionStage.QUARTERFINAL,
            format_type=BattleFormat.BATTLE_ROYALE,
            order=1,
            name="Batalla multi-clasificador demo",
            status=MatchStatus.PENDING,
        )
        multi_states = list(TeamCompetitionState.objects.filter(competition=school_competition).select_related("team").order_by("team__robot_name")[:4])
        for slot, state in enumerate(multi_states, start=1):
            CompetitionBattleEntry.objects.create(battle=multi_battle, team=state.team, slot_order=slot, origin_label="Demo multi-clasificador")
            state.current_stage = CompetitionStage.QUARTERFINAL
            state.current_battle = multi_battle
            state.current_status = ParticipantStatus.ACTIVE
            state.save(update_fields=["current_stage", "current_battle", "current_status", "updated_at"])
        if len(multi_states) >= 2:
            set_battle_qualifiers(multi_battle, [multi_states[0].team_id, multi_states[1].team_id], manual_override=False)

        # Manual phase proposal fixture: enough for the modal preview without materializing.
        eligible = list(TeamCompetitionState.objects.filter(competition=school_competition).select_related("team").order_by("team__robot_name")[:6])
        manual_groups = [[state.team_id for state in chunk] for chunk in [eligible[:3], eligible[3:6]]]
        configuration = dict(school_competition.configuration or {})
        configuration["manual_phase_proposal"] = {
            "name": "Fase manual demo",
            "source_stage": CompetitionStage.GROUPS,
            "operation": "normal",
            "participant_ids": sorted(state.team_id for state in eligible),
            "participant_signature": "demo-participants-signature",
            "group_count": 2,
            "group_sizes": [len(group) for group in manual_groups],
            "groups": manual_groups,
            "distribution": "balanced",
            "qualifier_targets": [1, 1],
            "qualifiers_per_group": 1,
            "total_qualifiers": 2,
            "distribution_label": "Balanceada",
            "seed": 20260724,
            "proposal_signature": "demo-manual-phase-signature",
        }
        school_competition.configuration = configuration
        school_competition.save(update_fields=["configuration", "updated_at"])

        # Communications fixtures.
        template_file = create_docx_template("invitacion_demo_banco.docx")
        template = CommunicationTemplate.objects.create(
            name="Plantilla invitacion demo banco",
            template_type="colegio",
            file=template_file,
            required_markers=["{{NOMBRE}}", "{{INSTITUCION}}", "{{EVENTO}}"],
            description="Plantilla ficticia para capturas base.",
            is_active=True,
        )
        recipients = []
        for idx, recipient_type in enumerate(["colegio", "universidad", "patrocinador", "colegio"], start=1):
            recipients.append(CommunicationRecipient.objects.create(
                name=f"Destinatario Demo {idx}",
                email=f"destinatario{idx}@example.com",
                recipient_type=recipient_type,
                institution_name=f"Institucion Comunicacion Demo {idx}",
                extra_data={"EVENTO": "Torneo Demo Manual de Usuario"},
            ))
        batch = CommunicationBatch.objects.create(
            name="Lote demo dry-run banco",
            communication_type="colegio",
            template=template,
            status=CommunicationStatus.PREVIEWED,
            send_mode=CommunicationSendMode.DRY_RUN,
            dry_run=True,
            created_by=user,
        )
        for recipient in recipients[:3]:
            CommunicationLog.objects.create(
                batch=batch,
                recipient_name=recipient.name,
                recipient_email=recipient.email,
                original_recipient_email=recipient.email,
                physical_recipient_email=recipient.email,
                communication_type=recipient.recipient_type,
                subject="Invitacion demo",
                status=CommunicationLogStatus.DRY_RUN,
                rendered_preview=f"NOMBRE: {recipient.name}",
            )

        # Generic phase detail fixture.
        # TournamentPhase/Match omitted intentionally: the bank focuses on active panel states.

    summary = {
        "database": str(connection.settings_dict["NAME"]),
        "users": User.objects.count(),
        "teams": Team.objects.count(),
        "school_registrations": SchoolRegistration.objects.count(),
        "university_registrations": UniversityRegistration.objects.count(),
        "competition_count": DivisionCompetition.objects.count(),
        "battles": CompetitionBattle.objects.count(),
        "communication_recipients": CommunicationRecipient.objects.count(),
        "communication_batches": CommunicationBatch.objects.count(),
        "attendance_tokens": {
            "pending_school": str(pending_school.attendance_token),
            "confirmed_school": str(confirmed_school.attendance_token),
        },
        "competitions": {
            "school": school_competition.id,
            "university": university_competition.id,
        },
        "advanced_states": [
            "grupo pendiente",
            "grupo con clasificados",
            "batalla con ganador",
            "batalla con clasificados multiples",
            "repechaje configurado",
            "fase progresiva",
            "fase manual propuesta",
            "podio guardado",
        ],
    }
    summary_path = settings.DEMO_ROOT / "datos" / "resumen_banco_capturas.json"
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
