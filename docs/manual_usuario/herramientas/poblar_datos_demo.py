import json
import os
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "docs.manual_usuario.herramientas.demo_settings")

import django  # noqa: E402
from django.conf import settings  # noqa: E402
from django.contrib.auth import get_user_model  # noqa: E402
from django.db import connection, transaction  # noqa: E402
from docx import Document  # noqa: E402


def assert_demo_database() -> None:
    db_path = Path(settings.DATABASES["default"]["NAME"]).resolve()
    demo_root = settings.DEMO_ROOT.resolve()
    official_db = (settings.BASE_DIR / "db.sqlite3").resolve()
    if demo_root not in db_path.parents:
        raise SystemExit(f"La base no esta dentro del entorno demo: {db_path}")
    if db_path == official_db:
        raise SystemExit("La base demo apunta a la base oficial. Abortado.")


def create_docx_template(relative_name: str) -> str:
    target_dir = Path(settings.MEDIA_ROOT) / "communication_templates"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / relative_name
    doc = Document()
    doc.add_heading("Invitacion demo Robot Explota Globos", level=1)
    doc.add_paragraph("Hola {{NOMBRE}}, invitamos a {{INSTITUCION}} al {{EVENTO}}.")
    doc.add_paragraph("Contacto ficticio: {{CORREO}}")
    doc.save(target)
    return f"communication_templates/{relative_name}"


def main() -> None:
    django.setup()
    assert_demo_database()

    from apps.participants.models import (  # noqa: E402
        Institution,
        InstitutionType,
        RegistrationStatus,
        SchoolParticipant,
        SchoolRegistration,
        Team,
        TeamMember,
        UniversityParticipant,
        UniversityRegistration,
    )
    from apps.tournament.models import (  # noqa: E402
        CompetitionBattle,
        CompetitionBattleEntry,
        CompetitionHistoryEntry,
        DivisionCompetition,
        DivisionGroup,
        DivisionGroupEntry,
        RuleSection,
        TeamCompetitionState,
        TournamentEdition,
    )
    from apps.tournament.services import initialize_competition, set_group_qualifiers  # noqa: E402
    from apps.invitations.models import (  # noqa: E402
        CommunicationBatch,
        CommunicationLog,
        CommunicationLogStatus,
        CommunicationRecipient,
        CommunicationSendMode,
        CommunicationStatus,
        CommunicationTemplate,
    )

    User = get_user_model()
    with transaction.atomic():
        User.objects.filter(username="operador_documentacion").delete()
        user = User.objects.create_user(
            username="operador_documentacion",
            email="operador@example.com",
            password="DemoManualUsuario2026!",
            first_name="Operador",
            last_name="Documentacion",
            is_staff=True,
            is_superuser=True,
        )

        CommunicationLog.objects.all().delete()
        CommunicationBatch.objects.all().delete()
        CommunicationRecipient.objects.all().delete()
        CommunicationTemplate.objects.all().delete()
        CompetitionHistoryEntry.objects.all().delete()
        TeamCompetitionState.objects.all().delete()
        CompetitionBattleEntry.objects.all().delete()
        CompetitionBattle.objects.all().delete()
        DivisionGroupEntry.objects.all().delete()
        DivisionGroup.objects.all().delete()
        DivisionCompetition.objects.all().delete()
        TeamMember.objects.all().delete()
        Team.objects.all().delete()
        Institution.objects.all().delete()
        SchoolParticipant.objects.all().delete()
        UniversityParticipant.objects.all().delete()
        SchoolRegistration.objects.all().delete()
        UniversityRegistration.objects.all().delete()
        RuleSection.objects.all().delete()
        TournamentEdition.objects.all().delete()

        edition = TournamentEdition.objects.create(
            name="Edicion Demo Manual de Usuario",
            description="Edicion ficticia y aislada para capturas de documentacion visual.",
            location="Laboratorio Demo UTP",
            is_active=True,
        )
        for order, title in enumerate(["Registro demo", "Fase de grupos", "Comunicaciones seguras"], start=1):
            RuleSection.objects.create(
                edition=edition,
                title=title,
                summary=f"Resumen ficticio para {title.lower()}.",
                body=f"Contenido demostrativo de {title.lower()} sin datos reales.",
                order=order,
                is_published=True,
            )

        school_inst = Institution.objects.create(name="Institucion Demo Norte", institution_type=InstitutionType.SCHOOL, city="Ciudad Demo")
        school_inst_2 = Institution.objects.create(name="Colegio Demo Sur", institution_type=InstitutionType.SCHOOL, city="Ciudad Demo")
        uni_inst = Institution.objects.create(name="Universidad de Ejemplo", institution_type=InstitutionType.UNIVERSITY, city="Ciudad Demo")

        team_specs = [
            ("Equipo Andromeda", school_inst, "Colegios"),
            ("Equipo Vector", school_inst_2, "Colegios"),
            ("Equipo Kilobyte", school_inst, "Colegios"),
            ("Equipo Prisma", school_inst_2, "Colegios"),
            ("Equipo Nebula", school_inst, "Colegios"),
            ("Equipo Pixel", school_inst_2, "Colegios"),
            ("Equipo Quasar", uni_inst, "Universidades"),
            ("Equipo Lambda", uni_inst, "Universidades"),
            ("Equipo Sigma", uni_inst, "Universidades"),
            ("Equipo Delta", uni_inst, "Universidades"),
        ]
        for index, (name, institution, category) in enumerate(team_specs, start=1):
            team = Team.objects.create(
                edition=edition,
                name=name,
                institution=institution,
                category_label=category,
                robot_name=name.replace("Equipo ", "Robot "),
                coach_name=f"Responsable Demo {index:02d}",
                coach_email=f"responsable{index:02d}@example.com",
                coach_phone=f"3000000{index:03d}",
                status=RegistrationStatus.APPROVED,
                notes="Equipo ficticio para capturas del Manual de Usuario.",
            )
            TeamMember.objects.create(
                team=team,
                full_name=f"Participante Demo {index:02d}",
                document_number=f"900000{index:04d}",
                email=f"participante{index:02d}@example.com",
                age=18 if category == "Universidades" else 15,
                role="lider",
                is_team_lead=True,
            )

        for index in range(1, 3):
            reg = SchoolRegistration.objects.create(
                edition=edition,
                institution_name=f"Institucion Registro Demo {index}",
                responsible_name=f"Responsable Registro Demo {index}",
                robot_name=f"Robot Registro Demo {index}",
                contact_phone=f"3110000{index:03d}",
                contact_email=f"registro{index}@example.com",
                status=RegistrationStatus.SUBMITTED,
            )
            SchoolParticipant.objects.create(
                registration=reg,
                full_name=f"Participante Registro Demo {index}",
                document_number=f"910000{index:04d}",
                role="lider",
                is_team_lead=True,
            )

        ureg = UniversityRegistration.objects.create(
            edition=edition,
            institution_name="Universidad Registro Demo",
            responsible_name="Responsable Universidad Demo",
            robot_name="Robot Universidad Demo",
            semester=3,
            contact_phone="3120000001",
            contact_email="universidad.demo@example.com",
            status=RegistrationStatus.SUBMITTED,
        )
        UniversityParticipant.objects.create(
            registration=ureg,
            full_name="Participante Universidad Demo",
            document_number="9200000001",
            role="lider",
            is_team_lead=True,
        )

        school_competition = initialize_competition(edition, "school")
        university_competition = initialize_competition(edition, "university")
        first_group = school_competition.groups.prefetch_related("entries__team").order_by("order").first()
        if first_group:
            selected = [str(entry.id) for entry in first_group.entries.order_by("slot_order")[:2]]
            if selected:
                set_group_qualifiers(first_group, selected, manual_override=False)

        template_file = create_docx_template("invitacion_demo.docx")
        template = CommunicationTemplate.objects.create(
            name="Plantilla invitacion demo",
            template_type="colegio",
            file=template_file,
            required_markers=["{{NOMBRE}}", "{{INSTITUCION}}", "{{EVENTO}}"],
            description="Plantilla ficticia para pruebas de documentacion.",
            is_active=True,
        )
        for idx, recipient_type in enumerate(["colegio", "universidad", "patrocinador"], start=1):
            CommunicationRecipient.objects.create(
                name=f"Destinatario Demo {idx}",
                email=f"destinatario{idx}@example.com",
                recipient_type=recipient_type,
                institution_name=f"Institucion Comunicacion Demo {idx}",
                extra_data={"EVENTO": "Torneo Demo Manual de Usuario"},
            )
        batch = CommunicationBatch.objects.create(
            name="Lote demo simulado",
            communication_type="colegio",
            template=template,
            status=CommunicationStatus.PREVIEWED,
            send_mode=CommunicationSendMode.DRY_RUN,
            dry_run=True,
            created_by=user,
        )
        CommunicationLog.objects.create(
            batch=batch,
            recipient_name="Destinatario Demo 1",
            recipient_email="destinatario1@example.com",
            original_recipient_email="destinatario1@example.com",
            physical_recipient_email="destinatario1@example.com",
            communication_type="colegio",
            subject="Invitacion demo",
            status=CommunicationLogStatus.DRY_RUN,
            rendered_preview="NOMBRE: Destinatario Demo 1",
        )

    summary = {
        "database": str(connection.settings_dict["NAME"]),
        "users": User.objects.count(),
        "teams": Team.objects.count(),
        "school_registrations": SchoolRegistration.objects.count(),
        "university_registrations": UniversityRegistration.objects.count(),
        "competitions": 2,
        "communication_recipients": CommunicationRecipient.objects.count(),
        "communication_batches": CommunicationBatch.objects.count(),
    }
    summary_path = settings.DEMO_ROOT / "datos" / "resumen_datos_demo.json"
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
